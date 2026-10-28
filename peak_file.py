"""
peak_file.py - The "Save peak info" Excel file: the columns a saved row can
have, reading and writing the .xlsx files, and the list of the files saved
to lately.

A peak file has one sheet named "Peaks". Row 1 holds the column headers, and
each row below it is one saved peak. Columns are found by their header, so
their order doesn't matter, and a header the app doesn't know is a text
column. The header text is part of the file format: once files exist, never
rename one, or old files lose that column. Adding items is fine.

Values are saved as the dialog shows them: rounded to the column's decimals,
with Excel's number format set to the same decimals.

Rows are never overwritten. Saving the same peak again (see SAME_PEAK_KEYS)
adds a row like any other and marks the earlier one replaced: its Status
cell, and grey text. So a sample's history stays in the file.

No Dash here. plot.py and peak_dialog.py pass in plain values.
"""
import datetime as dt
import os
import unicodedata
from copy import copy
from typing import NamedTuple, Optional

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SHEET = 'Peaks'
# A new row replaces a current saved row with the same values here and a λ
# this close (nm): the same peak, measured on the same day with the same
# instruments. A file without one of these columns isn't compared on it.
SAME_PEAK_KEYS = ['label', 'date', 'source', 'detector']
SAME_PEAK_NM = 0.01
# The Status column's values. Any text other than REPLACED counts as current.
CURRENT, REPLACED = 'current', 'replaced'
REPLACED_COLOR = 'FF808080'  # a replaced row's text
MAX_DECIMALS = 10           # the same limit as the tables' Customize dialog
DATE_FORMAT = 'mm/dd/yyyy'  # Excel's format for the Date column
TIME_FORMAT = 'hh:mm:ss'    # Excel's format for the Time column
TEXT = 'text'               # the column dropdown's value for "Text column…"


class PeakFileError(Exception):
    """A problem with a peak file, worded for the dialog."""


class Item(NamedTuple):
    key: str                        # stable id, used by the dialog and its session memory
    header: str                     # the column header, in the file and in the dropdown
    group: str                      # 'Fixed' columns are always there; the others are picked
    decimals: Optional[int] = None  # starting decimals; None = no decimals control


ITEMS = [
    Item('date',         'Date',               'Fixed'),
    Item('time',         'Time',               'Fixed'),
    Item('source',       'Source',             'Fixed'),
    Item('detector',     'Detector',           'Fixed'),
    Item('label',        'Label (SN)',         'Fixed'),
    Item('wl',           'λ (nm)',             'Fixed', 4),
    Item('status',       'Status',             'Fixed'),
    Item('depth_max',    'Depth, max (dB)',    'Peak', 3),
    Item('depth_avg',    'Depth, avg (dB)',    'Peak', 3),
    Item('depth_min',    'Depth, min (dB)',    'Peak', 3),
    Item('depth_custom', 'Depth, custom (dB)', 'Peak', 3),
    Item('width_max',    'Width, max (pm)',    'Peak', 1),
    Item('width_avg',    'Width, avg (pm)',    'Peak', 1),
    Item('width_min',    'Width, min (pm)',    'Peak', 1),
    Item('width_custom', 'Width, custom (pm)', 'Peak', 1),
    Item('il',           'IL (dB)',            'Peak', 1),
    Item('il_custom',    'IL, custom (dB)',    'Peak', 1),
    Item('wl_start',     'λ start (nm)',       'Sweep'),
    Item('wl_stop',      'λ stop (nm)',        'Sweep'),
    Item('speed',        'Sweep speed (nm/s)', 'Sweep'),
    Item('step',         'Step size (pm)',     'Sweep'),
    Item('tls_power',    'TLS power (dBm)',    'Sweep'),
    Item('avg_time',     'Averaging time (µs)', 'Sweep'),
    Item('pm_range',     'Power range (dBm)',  'Sweep'),
    Item('dyn_scans',    'Scans',              'Sweep'),
    Item('decrement',    'Decrement (dB)',     'Sweep'),
    Item('temperature',  'Temperature (°C)',   'Run'),
]
ITEM = {i.key: i for i in ITEMS}
# A new file's columns always start with the fixed items. An existing file
# only needs the ones the app relies on: Date, Label (SN) and λ, to find the
# row a new one replaces (see SAME_PEAK_KEYS), and Status, to mark it.
FIXED = [i.key for i in ITEMS if i.group == 'Fixed']
REQUIRED = ['date', 'label', 'wl', 'status']
RECENT_ROWS = 5    # the saved rows the dialog shows from an existing file
RECENT_FILES = 10  # the files the recent list keeps (see remember)
DEFAULT_DECIMALS = {i.key: i.decimals for i in ITEMS if i.decimals is not None}


def norm(header):
    """A header as compared: Unicode-normalized (so µ and μ match), with case
    and extra spaces ignored."""
    return ' '.join(unicodedata.normalize('NFKC', str(header)).casefold().split())


_KEY_BY_HEADER = {norm(i.header): i.key for i in ITEMS}


def column_options():
    """The column dropdown's list: the items that aren't fixed, under greyed
    group headings that can't be picked, then "Text column…"."""
    opts = []
    for group in ('Peak', 'Sweep', 'Run'):
        opts.append({'label': f'——— {group} ———', 'value': f'__{group}', 'disabled': True})
        opts += [{'label': i.header, 'value': i.key} for i in ITEMS if i.group == group]
    opts.append({'label': '——— Other ———', 'value': '__other', 'disabled': True})
    opts.append({'label': 'Text column…', 'value': TEXT})
    return opts


def column_problem(columns):
    """Why a new file's columns can't be saved, or None. columns: (item key,
    or None for a text column; header) pairs in order, fixed ones included."""
    seen = set()
    for key, header in columns:
        name = header.strip()
        if key is None:
            if not name:
                return 'Type a name for each text column.'
            if norm(name) in _KEY_BY_HEADER:
                return f'"{name}" is a built-in column. Give the text column another name.'
        if norm(name) in seen:
            return f'"{name}" is in the table twice. Remove one.'
        seen.add(norm(name))
    return None


# ---- values ---------------------------------------------------------------------
def parse_date(text):
    """A sweep's date (Params.date) as a date: MM/DD/YYYY since v1.1.6,
    YYYY-MM-DD before that. None if it's neither."""
    for fmt in ('%m/%d/%Y', '%Y-%m-%d'):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except (TypeError, ValueError):
            pass
    return None


def parse_time(text):
    """A sweep's time of day from Params.time, which Run stamps as
    YYYY-MM-DD_HH-MM-SS (config_helper.TIME_FORMAT), and stamped as
    MM/DD/YYYY_HH-MM-SS from v1.1.6 until 2026-09. None if it's neither."""
    for fmt in ('%Y-%m-%d_%H-%M-%S', '%m/%d/%Y_%H-%M-%S'):
        try:
            return dt.datetime.strptime(text, fmt).time()
        except (TypeError, ValueError):
            pass
    return None


def sweep_values(params, il, r):
    """The values every row of one sweep's range r shares, by item key. il:
    the range's insertion loss, or None without a reference."""
    return {
        'date': parse_date(params.date),
        'time': parse_time(params.time),
        'il': il,
        'wl_start': params.wl_start[r],
        'wl_stop': params.wl_stop[r],
        'speed': params.speed,
        'step': params.step_pm,
        'tls_power': params.tls_dbm,
        'avg_time': params.at_us,
        'pm_range': params.pm_range,
        'dyn_scans': params.dyn_scans,
        'decrement': params.decrement,
        'source': params.source,
        'detector': params.detector,
    }


def _is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def decimals_of(key, decimals):
    """The decimals an item uses, or None for an item without a decimals control."""
    if key not in DEFAULT_DECIMALS:
        return None
    return decimals.get(key, DEFAULT_DECIMALS[key])


def shown(key, value, decimals):
    """A value as the dialog shows it, and as it's saved: an item with a
    decimals control rounded to them, anything else unchanged. None stays None."""
    d = decimals_of(key, decimals)
    if d is not None and _is_number(value):
        return float(f'{value:.{d}f}')
    return value


def display(key, value, decimals):
    """A value's text in the dialog's table. "—" when there's none."""
    if value is None or value == '':
        return '—'
    if isinstance(value, dt.date):
        return value.strftime('%m/%d/%Y')
    if isinstance(value, dt.time):
        return value.strftime('%H:%M:%S')
    d = decimals_of(key, decimals)
    if d is not None and _is_number(value):
        return f'{value:.{d}f}'
    if isinstance(value, float):
        return f'{value:.10g}'
    return str(value)


def number_format(decimals):
    """Excel's number format for a number with this many decimals."""
    return '0' if decimals == 0 else '0.' + '0' * decimals


def format_decimals(fmt):
    """The decimals an Excel number format shows, like 3 for '0.000' or
    '#,##0.000'. None for General, text, dates and other formats."""
    if not fmt or fmt in ('General', '@'):
        return None
    section = fmt.split(';')[0]
    if '.' not in section:
        return 0 if ('0' in section or '#' in section) and '/' not in section else None
    n = 0
    for ch in section.split('.', 1)[1]:
        if ch not in '0#':
            break
        n += 1
    return n


# ---- the recent files and the files themselves -----------------------------------
def same_file(a, b):
    """Whether two paths name the same file. Case doesn't count where the
    system ignores it (Windows)."""
    return os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b))


def recent_files(list_path):
    """The peak files saved to lately, newest first: the lines of the text
    file at list_path, one path per line. [] when it doesn't exist yet or
    can't be read."""
    try:
        with open(list_path, encoding='utf-8') as f:
            return [line.strip() for line in f if line.strip()]
    except (OSError, ValueError):   # ValueError: not UTF-8
        return []


def remember(list_path, path):
    """Put path at the top of the recent list at list_path, which keeps the
    newest RECENT_FILES. A file that's gone stays on the list, in case its
    drive comes back, until newer files push it off. Raises OSError when the
    list can't be written."""
    path = os.path.abspath(path)
    paths = [path] + [p for p in recent_files(list_path) if not same_file(p, path)]
    tmp = f'{list_path}.tmp'
    try:
        with open(tmp, 'w', encoding='utf-8') as f:
            f.writelines(p + '\n' for p in paths[:RECENT_FILES])
        os.replace(tmp, list_path)
    except OSError:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise


def in_use(path):
    """True when another program has the file open for writing. On Windows,
    Excel locks a workbook it has open, so this is how a save would fail."""
    try:
        with open(path, 'r+b'):
            return False
    except PermissionError:
        return True
    except OSError:
        return False


def _sheet(wb):
    return wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]


def _load(path):
    name = os.path.basename(path)
    try:
        return openpyxl.load_workbook(path)
    except PermissionError:
        raise PeakFileError(f"Can't open {name}. If another program is using it, close it.")
    except FileNotFoundError:
        raise PeakFileError(f"{name} isn't there anymore. It may have been moved or deleted.")
    except Exception as e:
        raise PeakFileError(f"Can't read {name}: it isn't a valid .xlsx file ({e}).")


def _last_row(ws):
    """The last row with a value in it; 1 when there's only the header."""
    for r in range(ws.max_row, 1, -1):
        if any(c.value not in (None, '') for c in ws[r]):
            return r
    return 1


def _layout(ws, name):
    """A sheet's columns as (column number, item key or None, header), and the
    decimals its numbers show (item key -> decimals, from the last row that
    has a number in that column)."""
    headers = [c.value for c in ws[1]]
    while headers and (headers[-1] is None or str(headers[-1]).strip() == ''):
        headers.pop()
    if not headers:
        raise PeakFileError(f'{name} has no column headers in row 1.')
    cols, seen = [], set()
    for idx, h in enumerate(headers, start=1):
        text = '' if h is None else str(h).strip()
        if not text:
            raise PeakFileError(f'Column {get_column_letter(idx)} of {name} has no header.')
        if norm(text) in seen:
            raise PeakFileError(f'{name} has two "{text}" columns.')
        seen.add(norm(text))
        cols.append((idx, _KEY_BY_HEADER.get(norm(text)), text))
    keys = {key for _, key, _ in cols}
    missing = [ITEM[k].header for k in REQUIRED if k not in keys]
    if missing:
        raise PeakFileError(f'{name} has no {" or ".join(missing)} column, so it '
                            "can't be used for peaks.")
    last = _last_row(ws)
    decimals = {}
    for idx, key, _ in cols:
        if key in DEFAULT_DECIMALS:
            for r in range(last, 1, -1):
                cell = ws.cell(r, idx)
                if _is_number(cell.value):
                    d = format_decimals(cell.number_format)
                    if d is not None:
                        decimals[key] = min(d, MAX_DECIMALS)
                    break
    return cols, decimals


def _match_value(key, value):
    """A value as the same-peak rule compares it (SAME_PEAK_KEYS): a date as
    YYYY-MM-DD, text without its outer spaces, None when it's empty. Plain
    data, since the dialog keeps the saved rows in a dcc.Store."""
    if isinstance(value, dt.date):   # a datetime too
        return value.strftime('%Y-%m-%d')
    text = '' if value is None else str(value).strip()
    if key == 'date' and text:
        date = parse_date(text)
        return date.isoformat() if date else text
    return text or None


def _is_replaced(status):
    """Whether a Status cell marks its row replaced."""
    return str(status or '').strip().casefold() == REPLACED


def _saved(ws, cols):
    """The current rows a new row can replace: {'row': Excel row, 'match':
    {key: value}, 'wl': λ} for each row with a label and a λ that isn't
    marked replaced. 'match' holds the row's SAME_PEAK_KEYS values that the
    file has columns for, as _match_value gives them."""
    col = {key: idx for idx, key, _ in cols if key}
    keys = [k for k in SAME_PEAK_KEYS if k in col]
    out = []
    for r in range(2, _last_row(ws) + 1):
        if _is_replaced(ws.cell(r, col['status']).value):
            continue
        match = {k: _match_value(k, ws.cell(r, col[k]).value) for k in keys}
        wl = ws.cell(r, col['wl']).value
        if match['label'] and _is_number(wl):
            out.append({'row': r, 'match': match, 'wl': float(wl)})
    return out


def _same_sweep(match, sweep):
    """Whether a saved row's 'match' values, the label aside, are the
    sweep's (item key -> value): only then can its peaks replace the row."""
    return all(_match_value(k, sweep.get(k)) == v for k, v in match.items() if k != 'label')


def _cell_text(cell):
    """A saved cell as the file shows it: a number with its cell's decimals,
    a date as MM/DD/YYYY, a time as HH:MM:SS. "—" when it's empty."""
    v = cell.value
    if v is None or v == '':
        return '—'
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime('%m/%d/%Y')
    if isinstance(v, dt.time):
        return v.strftime('%H:%M:%S')
    if _is_number(v):
        d = format_decimals(cell.number_format)
        if d is not None:
            return f'{v:.{d}f}'
        return f'{v:.10g}' if isinstance(v, float) else str(v)
    return str(v)


def _recent(ws, cols):
    """The last RECENT_ROWS saved rows, top to bottom: {'row': Excel row,
    'cells': [text per column]}."""
    out = []
    for r in range(_last_row(ws), 1, -1):
        cells = [ws.cell(r, idx) for idx, _, _ in cols]
        if any(c.value not in (None, '') for c in cells):
            out.append({'row': r, 'cells': [_cell_text(c) for c in cells]})
            if len(out) == RECENT_ROWS:
                break
    return out[::-1]


def read(path, sweep=None):
    """What the dialog needs from a peak file: its columns in order ({'item':
    key} or {'text': header}), the decimals its numbers show, the last few
    rows as text (see _recent), and the current rows that this sweep's peaks
    could replace (see _saved), each with its 'cells' as text too.
    sweep: the sweep's values (item key -> value), whose date, source and
    detector those rows share; None for every current row.
    Raises PeakFileError."""
    wb = _load(path)
    try:
        ws = _sheet(wb)
        cols, decimals = _layout(ws, os.path.basename(path))
        saved = [s for s in _saved(ws, cols) if sweep is None or _same_sweep(s['match'], sweep)]
        for s in saved:
            s['cells'] = [_cell_text(ws.cell(s['row'], idx)) for idx, _, _ in cols]
        return {'columns': [{'item': key} if key else {'text': header} for _, key, header in cols],
                'decimals': decimals,
                'saved': saved,
                'recent': _recent(ws, cols)}
    finally:
        wb.close()


def find_saved(saved, values):
    """The Excel row a new row with these values (item key -> value) replaces:
    the current saved row (see _saved) with the same SAME_PEAK_KEYS values and
    the closest λ within SAME_PEAK_NM. None if none, or if the new row has no λ."""
    wl = values.get('wl')
    if not _is_number(wl):
        return None
    best = None
    for s in saved:
        gap = abs(s['wl'] - wl)
        if (gap <= SAME_PEAK_NM and (best is None or gap < best[1])
                and all(_match_value(k, values.get(k)) == v for k, v in s['match'].items())):
            best = (s['row'], gap)
    return best[0] if best else None


# ---- writing ------------------------------------------------------------------------
def _write(cell, key, value, decimals):
    """One cell, as the dialog showed it. key None = a text column."""
    if value is None or value == '':
        cell.value = None
    elif key is None or key == 'label':
        cell.value = str(value)
        cell.number_format = '@'   # so Excel keeps a typed 0160060 as text too
    elif key == 'date':
        cell.value = value
        cell.number_format = DATE_FORMAT
    elif key == 'time':
        cell.value = value
        cell.number_format = TIME_FORMAT
    else:
        cell.value = shown(key, value, decimals)
        d = decimals_of(key, decimals)
        if d is not None and _is_number(value):
            cell.number_format = number_format(d)


def _mark_replaced(ws, row, cols):
    """Mark a saved row replaced: REPLACED in its Status cell, and grey text."""
    for idx, key, _ in cols:
        cell = ws.cell(row, idx)
        if key == 'status':
            cell.value = REPLACED
        font = copy(cell.font)
        font.color = REPLACED_COLOR
        cell.font = font


def _save_atomic(wb, path):
    """Save to a temporary file, then put it in place, so a crash can't leave
    half a file. On Windows, replacing a file Excel has open fails."""
    name = os.path.basename(path)
    tmp = path + '.tmp'
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except OSError as e:
        try:
            os.remove(tmp)
        except OSError:
            pass
        if isinstance(e, PermissionError):
            raise PeakFileError(f"Can't save {name}. If it's open in Excel, "
                                'close it, then press Save again.')
        raise PeakFileError(f"Can't save {name}: {e.strerror or e}.")


def save(path, rows, *, columns=None, decimals=None):
    """Add rows to the peak file at path, after its last row, and return
    (added, replaced): the Excel rows written, and the saved rows they
    replaced.

    rows: one per peak, item key -> value, plus 'texts': {header: text} for
    the text columns. Values are rounded to the decimals here. Their Status
    is CURRENT.
    columns: for a new file, its columns as (item key or None, header) pairs
    in order. Without it, the file at path must exist, and its header sets
    the columns and its numbers set the decimals.
    decimals: item key -> decimals for a new file, and the fallback for an
    existing file's column that has no number yet.
    A row that matches a current saved row (find_saved: the same label, date,
    source and detector, and a λ within SAME_PEAK_NM) marks that row
    replaced (see _mark_replaced); nothing else in it changes. Raises
    PeakFileError.
    """
    name = os.path.basename(path)
    decimals = {**DEFAULT_DECIMALS, **(decimals or {})}
    if columns is not None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = SHEET
        cols = [(i, key, header) for i, (key, header) in enumerate(columns, start=1)]
        bold = Font(bold=True)
        for idx, _, header in cols:
            cell = ws.cell(1, idx, header)
            cell.font = bold
            ws.column_dimensions[get_column_letter(idx)].width = max(11, len(header) + 3)
        ws.freeze_panes = 'A2'
        saved, next_row = [], 2
    else:
        wb = _load(path)
        ws = _sheet(wb)
        cols, file_decimals = _layout(ws, name)
        decimals.update(file_decimals)
        saved, next_row = _saved(ws, cols), _last_row(ws) + 1

    added, replaced = [], []
    for values in ({**v, 'status': CURRENT} for v in rows):
        texts = {norm(h): t for h, t in values.get('texts', {}).items()}
        old = find_saved(saved, values)
        if old is not None:
            _mark_replaced(ws, old, cols)
            saved = [s for s in saved if s['row'] != old]
            replaced.append(old)
        added.append(next_row)
        for idx, key, header in cols:
            value = values.get(key) if key else texts.get(norm(header))
            _write(ws.cell(next_row, idx), key, value, decimals)
        next_row += 1
    # Excel's filter buttons on the header row, stretched over every column
    # and row, new ones included. A filter set in Excel keeps its first cell,
    # so its settings stay on their columns.
    first = ws.auto_filter.ref.split(':')[0] if ws.auto_filter.ref else 'A1'
    ws.auto_filter.ref = f'{first}:{get_column_letter(cols[-1][0])}{_last_row(ws)}'
    _save_atomic(wb, path)
    return added, replaced
